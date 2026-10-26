from openpyxl import load_workbook
path = "data/resultat_2025.xlsx"
wb = load_workbook(path)
print ("flikar före", wb.sheetnames)
for name in wb.sheetnames:
    if name != "Tabell 3":
        del wb[name]
wb.save(path)
print("flikar efter:", wb.sheetnames)
